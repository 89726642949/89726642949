from datetime import date
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from sistema_atividades import db


def _style_header_row(table: Table) -> None:
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    )


def generate_monthly_report(year: int, month: int, output_dir: Path) -> Tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    records = db.list_records_in_month(year, month)
    summary_category = db.get_counts_by_category(year, month)
    summary_status = db.get_counts_by_status(year, month)
    highlights = db.get_deadline_highlights(records)

    pdf_path = output_dir / f"relatorio_mensal_{year}_{month:02d}.pdf"
    xlsx_path = output_dir / f"relatorio_mensal_{year}_{month:02d}.xlsx"

    _build_monthly_pdf(pdf_path, year, month, records, summary_category, summary_status, highlights)
    _build_monthly_excel(xlsx_path, year, month, records, summary_category, summary_status, highlights)

    return pdf_path, xlsx_path


def _build_monthly_pdf(
    pdf_path: Path,
    year: int,
    month: int,
    records: List[Dict],
    summary_category: Dict[str, int],
    summary_status: Dict[str, int],
    highlights: List[Dict],
) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4)
    elements = []

    elements.append(Paragraph(f"Relatorio Mensal - {month:02d}/{year}", styles["Title"]))
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Resumo por categoria", styles["Heading2"]))
    cat_data = [["Categoria", "Total"]]
    for key, value in sorted(summary_category.items()):
        cat_data.append([key, str(value)])
    cat_table = Table(cat_data, hAlign="LEFT")
    _style_header_row(cat_table)
    elements.append(cat_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Resumo por status", styles["Heading2"]))
    status_data = [["Status", "Total"]]
    for key, value in sorted(summary_status.items()):
        status_data.append([key, str(value)])
    status_table = Table(status_data, hAlign="LEFT")
    _style_header_row(status_table)
    elements.append(status_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Registros do mes", styles["Heading2"]))
    reg_data = [["Data", "Categoria", "Numero", "Assunto", "Status"]]
    for rec in records:
        reg_data.append(
            [
                rec.get("record_date", ""),
                rec.get("category", ""),
                rec.get("doc_number", ""),
                rec.get("subject", ""),
                rec.get("status", ""),
            ]
        )
    reg_table = Table(reg_data, repeatRows=1)
    _style_header_row(reg_table)
    elements.append(reg_table)
    elements.append(Spacer(1, 12))

    if highlights:
        elements.append(Paragraph("Destaques de prazo", styles["Heading2"]))
        hl_data = [["Data", "Numero", "Assunto", "Prazo", "Status prazo"]]
        for rec in highlights:
            hl_data.append(
                [
                    rec.get("record_date", ""),
                    rec.get("doc_number", ""),
                    rec.get("subject", ""),
                    rec.get("protocol_deadline_date", ""),
                    rec.get("deadline_status", ""),
                ]
            )
        hl_table = Table(hl_data, repeatRows=1)
        _style_header_row(hl_table)
        elements.append(hl_table)

    doc.build(elements)


def _build_monthly_excel(
    xlsx_path: Path,
    year: int,
    month: int,
    records: List[Dict],
    summary_category: Dict[str, int],
    summary_status: Dict[str, int],
    highlights: List[Dict],
) -> None:
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Resumo"
    ws_summary.append([f"Relatorio Mensal {month:02d}/{year}"])
    ws_summary.append([])
    ws_summary.append(["Resumo por categoria"])
    ws_summary.append(["Categoria", "Total"])
    for key, value in sorted(summary_category.items()):
        ws_summary.append([key, value])
    ws_summary.append([])
    ws_summary.append(["Resumo por status"])
    ws_summary.append(["Status", "Total"])
    for key, value in sorted(summary_status.items()):
        ws_summary.append([key, value])

    for cell in ws_summary[1]:
        cell.font = Font(bold=True, size=12)

    ws_records = wb.create_sheet("Registros")
    ws_records.append(["Data", "Categoria", "Numero", "Assunto", "Status"])
    for rec in records:
        ws_records.append(
            [
                rec.get("record_date", ""),
                rec.get("category", ""),
                rec.get("doc_number", ""),
                rec.get("subject", ""),
                rec.get("status", ""),
            ]
        )

    ws_highlights = wb.create_sheet("Destaques")
    ws_highlights.append(["Data", "Numero", "Assunto", "Prazo", "Status prazo"])
    for rec in highlights:
        ws_highlights.append(
            [
                rec.get("record_date", ""),
                rec.get("doc_number", ""),
                rec.get("subject", ""),
                rec.get("protocol_deadline_date", ""),
                rec.get("deadline_status", ""),
            ]
        )

    _autosize_columns(ws_summary)
    _autosize_columns(ws_records)
    _autosize_columns(ws_highlights)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def generate_annual_report(year: int, output_dir: Path) -> Tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    records = db.list_records_in_year(year)
    monthly_totals = db.get_annual_monthly_totals(year)
    category_totals = db.get_annual_category_totals(year)
    status_totals = db.get_annual_status_totals(year)

    pdf_path = output_dir / f"relatorio_anual_{year}.pdf"
    xlsx_path = output_dir / f"relatorio_anual_{year}.xlsx"

    _build_annual_pdf(pdf_path, year, records, monthly_totals, category_totals, status_totals)
    _build_annual_excel(xlsx_path, year, records, monthly_totals, category_totals, status_totals)

    return pdf_path, xlsx_path


def _build_annual_pdf(
    pdf_path: Path,
    year: int,
    records: List[Dict],
    monthly_totals: Dict[int, int],
    category_totals: Dict[str, int],
    status_totals: Dict[str, int],
) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4)
    elements = []

    elements.append(Paragraph(f"Relatorio Anual - {year}", styles["Title"]))
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Totais por mes", styles["Heading2"]))
    month_data = [["Mes", "Total"]]
    for month in range(1, 13):
        month_data.append([f"{month:02d}", str(monthly_totals.get(month, 0))])
    month_table = Table(month_data, hAlign="LEFT")
    _style_header_row(month_table)
    elements.append(month_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Totais por categoria", styles["Heading2"]))
    cat_data = [["Categoria", "Total"]]
    for key, value in sorted(category_totals.items()):
        cat_data.append([key, str(value)])
    cat_table = Table(cat_data, hAlign="LEFT")
    _style_header_row(cat_table)
    elements.append(cat_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Totais por status", styles["Heading2"]))
    status_data = [["Status", "Total"]]
    for key, value in sorted(status_totals.items()):
        status_data.append([key, str(value)])
    status_table = Table(status_data, hAlign="LEFT")
    _style_header_row(status_table)
    elements.append(status_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("Registros consolidados", styles["Heading2"]))
    reg_data = [["Data", "Categoria", "Numero", "Assunto", "Status"]]
    for rec in records:
        reg_data.append(
            [
                rec.get("record_date", ""),
                rec.get("category", ""),
                rec.get("doc_number", ""),
                rec.get("subject", ""),
                rec.get("status", ""),
            ]
        )
    reg_table = Table(reg_data, repeatRows=1)
    _style_header_row(reg_table)
    elements.append(reg_table)

    doc.build(elements)


def _build_annual_excel(
    xlsx_path: Path,
    year: int,
    records: List[Dict],
    monthly_totals: Dict[int, int],
    category_totals: Dict[str, int],
    status_totals: Dict[str, int],
) -> None:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Resumo"
    ws_summary.append([f"Relatorio Anual {year}"])
    ws_summary.append([])

    ws_summary.append(["Totais por mes"])
    ws_summary.append(["Mes", "Total"])
    for month in range(1, 13):
        ws_summary.append([f"{month:02d}", monthly_totals.get(month, 0)])
    ws_summary.append([])

    ws_summary.append(["Totais por categoria"])
    ws_summary.append(["Categoria", "Total"])
    for key, value in sorted(category_totals.items()):
        ws_summary.append([key, value])
    ws_summary.append([])

    ws_summary.append(["Totais por status"])
    ws_summary.append(["Status", "Total"])
    for key, value in sorted(status_totals.items()):
        ws_summary.append([key, value])

    for cell in ws_summary[1]:
        cell.font = Font(bold=True, size=12)

    ws_records = wb.create_sheet("Registros")
    ws_records.append(["Data", "Categoria", "Numero", "Assunto", "Status"])
    for rec in records:
        ws_records.append(
            [
                rec.get("record_date", ""),
                rec.get("category", ""),
                rec.get("doc_number", ""),
                rec.get("subject", ""),
                rec.get("status", ""),
            ]
        )

    _autosize_columns(ws_summary)
    _autosize_columns(ws_records)

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)


def _autosize_columns(ws) -> None:
    for col_cells in ws.columns:
        length = max(len(str(cell.value or "")) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(50, length + 2)
